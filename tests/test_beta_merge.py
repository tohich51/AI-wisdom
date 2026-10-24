"""Synthetic regressions for the public-main/current-runtime merge boundary."""

from uuid import uuid4
from zipfile import ZipFile

import httpx
import pytest
from docx import Document
from docx.oxml import OxmlElement
from ebooklib import epub

from kb.catalog.fetch_html import fetch_html_snapshot
from kb.document_parser import parse_document
from kb.jobs.ingestion import _claim_job


def test_current_ingestion_reads_inserted_docx_text_and_unique_merged_cells(tmp_path):
    document = Document()
    p = document.add_paragraph("Before ")
    insertion = OxmlElement("w:ins")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "visible insertion"
    run.append(text)
    insertion.append(run)
    p._p.append(insertion)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(1, 0)).text = "One origin"
    table.cell(0, 1).text = "Upper"
    table.cell(1, 1).text = "Lower"
    path = tmp_path / "synthetic.docx"
    document.save(path)
    rows = parse_document(path, "application/octet-stream")
    assert rows[0]["text"] == "Before visible insertion"
    cells = [r for r in rows if r["locator_kind"] == "docx_cell"]
    assert [(r["cell_row"], r["cell_column"], r["text"]) for r in cells] == [
        (1, 1, "One origin"),
        (1, 2, "Upper"),
        (2, 2, "Lower"),
    ]
    assert all(r["file_page"] is None for r in rows)


def test_current_ingestion_keeps_inline_epub_text_in_one_structural_block(tmp_path):
    book = epub.EpubBook()
    book.set_identifier("synthetic-inline")
    book.set_title("Synthetic")
    book.set_language("en")
    chapter = epub.EpubHtml(title="Chapter", file_name="one.xhtml", uid="one")
    chapter.content = "<h1>Chapter</h1><p>Before <b>bold</b> after.</p>"
    book.add_item(chapter)
    book.spine = [chapter]
    book.toc = [chapter]
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    path = tmp_path / "synthetic.epub"
    epub.write_epub(str(path), book)
    rows = parse_document(path, "application/epub+zip")
    assert [r["text"] for r in rows] == ["Chapter", "Before bold after."]
    assert [(r["spine"], r["paragraph"]) for r in rows] == [("one", 1), ("one", 2)]


def test_snapshot_default_transport_is_pinned_and_has_no_redirect_cookies(monkeypatch):
    seen = []

    def handle(self, request):
        seen.append(request)
        if len(seen) == 1:
            return httpx.Response(
                302,
                headers={"location": "/next", "set-cookie": "canary=1"},
                content=b"",
                request=request,
            )
        return httpx.Response(
            200, headers={"content-type": "text/html"}, content=b"<p>Synthetic</p>", request=request
        )

    monkeypatch.setattr(httpx.HTTPTransport, "handle_request", handle)
    monkeypatch.setenv("HTTPS_PROXY", "http://127.0.0.1:9")
    result = fetch_html_snapshot(
        "https://example.test/first", resolver=lambda host, port: ["93.184.216.34"]
    )
    assert [r.url.host for r in seen] == ["93.184.216.34"] * 2
    assert all(r.headers["host"] == "example.test" for r in seen)
    assert all(r.extensions["sni_hostname"] == "example.test" for r in seen)
    assert all("cookie" not in r.headers for r in seen)
    assert result.requested_urls == ("https://example.test/first", "https://example.test/next")
    assert result.final_url == "https://example.test/next"


def test_snapshot_stops_stream_before_consuming_oversized_body():
    from kb.catalog.fetch_html import SnapshotPolicy
    from kb.catalog.upload_fetch import FetchTooLarge

    consumed = []

    class Stream(httpx.SyncByteStream):
        def __iter__(self):
            for chunk in (b"<p>", b"1234567890", b"must not be read"):
                consumed.append(chunk)
                yield chunk

    transport = httpx.MockTransport(
        lambda req: httpx.Response(
            200, headers={"content-type": "text/html"}, stream=Stream(), request=req
        )
    )
    with pytest.raises(FetchTooLarge):
        fetch_html_snapshot(
            "https://example.test/",
            SnapshotPolicy(max_bytes=5),
            transport=transport,
            resolver=lambda host, port: ["93.184.216.34"],
        )
    assert consumed == [b"<p>", b"1234567890"]


def test_current_worker_does_not_reclaim_needs_ocr():
    from kb.access.policy import Principal

    actor = uuid4()

    class Connection:
        def execute(self, *_args):
            return self

        def fetchone(self):
            return {"submitted_by": actor, "status": "needs_ocr"}

    assert _claim_job(Connection(), uuid4(), Principal(actor, uuid4())) is None


def test_beta_version_is_consistent():
    import tomllib
    from pathlib import Path

    import kb

    project = tomllib.loads((Path(__file__).resolve().parents[1] / "pyproject.toml").read_text())
    assert project["project"]["version"] == kb.__version__ == "0.1.0b1"


def test_snapshot_stream_failure_is_sanitized_and_closed():
    from kb.catalog.upload_fetch import UrlRefused

    closed = []

    class Stream(httpx.SyncByteStream):
        def __iter__(self):
            yield b"<p>"
            raise httpx.ReadError("private peer detail")

        def close(self):
            closed.append(True)

    transport = httpx.MockTransport(
        lambda req: httpx.Response(
            200, headers={"content-type": "text/html"}, stream=Stream(), request=req
        )
    )
    with pytest.raises(UrlRefused) as refusal:
        fetch_html_snapshot(
            "https://example.test/",
            transport=transport,
            resolver=lambda host, port: ["93.184.216.34"],
        )
    assert refusal.value.reason == "fetch_failed"
    assert "private peer detail" not in str(refusal.value)
    assert closed == [True]


@pytest.mark.parametrize(
    ("markup", "expected"),
    [
        (
            "<div>Before<p>Nested paragraph</p>After nested paragraph</div>",
            ["Before", "Nested paragraph", "After nested paragraph"],
        ),
        (
            "Loose <b>inline</b> text<div>Container <em>inline</em></div>Tail",
            ["Loose inline text", "Container inline", "Tail"],
        ),
        (
            "<blockquote>Outer<p>Inner</p>Outer tail</blockquote>"
            "<div>Next</div><section>Last</section>",
            ["Outer", "Inner", "Outer tail", "Next", "Last"],
        ),
        (
            "<div>Before<br/>after<script>hidden</script><style>hidden</style>"
            "<noscript>hidden</noscript><iframe>hidden</iframe> end</div>",
            ["Before after end"],
        ),
    ],
)
def test_current_ingestion_preserves_epub_container_text_and_provenance(tmp_path, markup, expected):
    book = epub.EpubBook()
    book.set_identifier("synthetic-container-text")
    book.set_title("Synthetic")
    book.set_language("en")
    chapters = []
    for ident, content in [("one", markup), ("two", "<p>Second spine</p>")]:
        chapter = epub.EpubHtml(title=ident, file_name=f"{ident}.xhtml", uid=ident)
        chapter.content = content
        book.add_item(chapter)
        chapters.append(chapter)
    book.spine = chapters
    book.toc = chapters
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    path = tmp_path / "synthetic-containers.epub"
    epub.write_epub(str(path), book)
    # EbookLib's writer drops direct body text. Put the exact synthetic XHTML
    # into the archive, like a book produced by another EPUB authoring tool.
    with ZipFile(path) as archive:
        entries = [(entry, archive.read(entry)) for entry in archive.infolist()]
    with ZipFile(path, "w") as archive:
        for entry, content in entries:
            if entry.filename.endswith("/one.xhtml"):
                content = (
                    '<html xmlns="http://www.w3.org/1999/xhtml"><head><title>one</title>'
                    f"</head><body>{markup}</body></html>"
                ).encode()
            archive.writestr(entry, content)
    rows = parse_document(path, "application/epub+zip")
    assert [r["text"] for r in rows] == [*expected, "Second spine"]
    assert [(r["spine"], r["paragraph"]) for r in rows] == [
        *(("one", n) for n in range(1, len(expected) + 1)),
        ("two", 1),
    ]
    assert [r["ordinal"] for r in rows] == list(range(len(rows)))
    assert all(r["locator_kind"] == "epub_paragraph" for r in rows)
    assert all(r["file_page"] is None and r["printed_label"] is None for r in rows)
    assert all(r["parser_name"] == "ebooklib" and r["parser_version"] for r in rows)


@pytest.mark.parametrize(
    ("routed_addresses", "reason"),
    [
        (["127.0.0.1"], "address_not_allowed"),
        (["::1"], "address_not_allowed"),
        (["93.184.216.34", "10.0.0.1"], "address_not_allowed"),
        (["not-an-ip"], "address_not_allowed"),
        ([], "dns_failure"),
        (OSError("private resolver detail"), "dns_failure"),
    ],
)
def test_snapshot_revalidates_actual_routed_addresses(monkeypatch, routed_addresses, reason):
    from kb.catalog.upload_fetch import UrlRefused

    answers = iter([["93.184.216.34"], ["93.184.216.34"], routed_addresses])
    resolved = []
    sent = []

    def resolve(host, port):
        resolved.append((host, port))
        answer = next(answers)
        if isinstance(answer, OSError):
            raise answer
        return answer

    def handle(self, request):
        sent.append(request)
        return httpx.Response(
            200, headers={"content-type": "text/html"}, content=b"<p>Synthetic</p>", request=request
        )

    # Real httpx Client and default HTTPTransport path; only socket I/O is
    # replaced. No private host is contacted even if this regression fails.
    monkeypatch.setattr(httpx.HTTPTransport, "handle_request", handle)
    with pytest.raises(UrlRefused) as refusal:
        fetch_html_snapshot("https://example.test/", resolver=resolve)
    assert refusal.value.reason == reason
    assert "private resolver detail" not in str(refusal.value)
    assert len(resolved) == 3
    assert sent == []
