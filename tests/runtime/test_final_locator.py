"""Real DOCX locations must survive rule retrieval and pinned MCP context."""

import hashlib
import io
import json
from uuid import uuid4

import psycopg
import pytest
from docx import Document
from test_runtime import rule
from test_runtime import runtime_db as runtime_db
from test_runtime import world as world

from kb.catalog import content, projects, upload
from kb.jobs.ingestion import run_job, submit_processing

pytestmark = pytest.mark.integration


def test_docx_cell_survives_rule_and_project_manifest(world):
    document = Document()
    document.add_paragraph("Synthetic marketing brief.")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "First table."
    table = document.add_table(rows=2, cols=2)
    for row in range(2):
        for column in range(2):
            table.cell(row, column).text = f"Segment {row + 1}, metric {column + 1}."
    payload = io.BytesIO()
    document.save(payload)
    payload.seek(0)

    principal = world["principal"]
    with psycopg.connect(world["dsn"]) as conn:
        submitted = upload.submit_file(
            conn,
            principal,
            world["store"],
            upload.SubmitFile(
                library_id=world["lib"], title="Synthetic brief", idempotency_key=str(uuid4())
            ),
            payload,
            declared_media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            filename="synthetic-brief.docx",
        )
        job = submit_processing(conn, principal, submitted.source.id, str(uuid4()), extract=False)
    run_job(world["worker"], job["job_id"], principal, world["store"].root)

    with psycopg.connect(world["dsn"]) as conn:
        source = content.read_source(conn, principal, submitted.source.id)
        fragment = next(
            item for item in source["fragments"] if item["text"] == "Segment 2, metric 2."
        )
        assert fragment["locator_kind"] == "docx_cell"
        assert (fragment["table_index"], fragment["cell_row"], fragment["cell_column"]) == (
            2,
            2,
            2,
        )
        published = rule(world, fragment["id"])
        raw = content.get_rule(conn, principal, published["id"], published["version_no"])
        raw_locator = raw["provenance"][0]["locator"]
        assert tuple(raw_locator[key] for key in ("table_index", "cell_row", "cell_column")) == (
            2,
            2,
            2,
        )

        project = projects.create_project(
            conn,
            principal,
            projects.CreateProject(organisation_id=principal.account_id, name="Synthetic launch"),
        ).project.id
        projects.pin_rule(
            conn,
            principal,
            project,
            projects.PinRule(rule_id=published["id"], version_no=published["version_no"]),
        )
        context = projects.project_context(conn, principal, project)
        assert context.complete
        provenance = context.pins[0].rule.provenance[0]
        assert (provenance.source_id, provenance.fragment_id) == (
            submitted.source.id,
            fragment["id"],
        )
        locator = provenance.locator
        assert locator.kind == "docx_cell" and locator.table == 2 and locator.cell == (2, 2)
        assert locator.file_page is None and locator.paragraph is None
        result = context.to_mcp_result()
        encoded_locator = result.manifest["pins"][0]["rule"]["provenance"][0]["locator"]
        assert encoded_locator == {"kind": "docx_cell", "table": 2, "cell": [2, 2]}
        canonical = json.dumps(
            result.manifest, sort_keys=True, separators=(",", ":"), ensure_ascii=False
        )
        assert result.manifest_sha256 == hashlib.sha256(canonical.encode()).hexdigest()
