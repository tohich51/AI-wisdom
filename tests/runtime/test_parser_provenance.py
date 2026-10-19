"""Actual document structures and real PG roundtrip; no books or model calls."""

import importlib.metadata
import platform
from uuid import uuid4

import psycopg
import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from ebooklib import epub
from test_runtime import runtime_db as runtime_db
from test_runtime import source as make_source
from test_runtime import world as world

from kb.access.policy import AccessDenied
from kb.catalog import content
from kb.document_parser import parse_document
from kb.jobs.ingestion import _chunks, run_job, submit_processing

pytestmark = pytest.mark.integration


def test_docx_heading_inheritance_and_cell_addresses(tmp_path):
    document = Document()
    document.add_heading("Section A", level=1)
    document.add_paragraph("Body A")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).add_paragraph("Cell heading", style="Heading 2")
    document.add_paragraph("After table")
    custom = document.styles.add_style("Synthetic custom section", WD_STYLE_TYPE.PARAGRAPH)
    custom.base_style = document.styles["Heading 2"]
    document.add_paragraph("Section B", style=custom)
    document.add_paragraph("Body B")
    path = tmp_path / "headings.docx"
    document.save(path)
    rows = parse_document(
        path, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    paragraphs = {r["paragraph"]: r for r in rows if r["locator_kind"] == "docx_paragraph"}
    assert {n: r["text"] for n, r in paragraphs.items()} == {
        1: "Section A",
        2: "Body A",
        3: "After table",
        4: "Section B",
        5: "Body B",
    }
    assert [paragraphs[n]["heading"] for n in range(1, 6)] == [
        "Section A",
        "Section A",
        "Section A",
        "Section B",
        "Section B",
    ]
    cell = next(r for r in rows if r["locator_kind"] == "docx_cell")
    assert (cell["table_index"], cell["cell_row"], cell["cell_column"]) == (1, 1, 1)
    assert cell["text"] == table.cell(0, 0).text and cell["heading"] == "Section A"
    assert all(
        r["parser_name"] == "python-docx"
        and r["parser_version"] == importlib.metadata.version("python-docx")
        for r in rows
    )
    assert all(r["file_page"] is None and r["spine"] is None for r in rows)
    chunks = list(_chunks([{**r, "id": uuid4()} for r in rows]))
    assert len(chunks) == 2
    assert [x["text"] for x in chunks[1]] == ["Section B", "Body B"]


def test_epub_repeated_heading_preserves_distinct_spine_items(tmp_path):
    book = epub.EpubBook()
    book.set_identifier("synthetic-provenance")
    book.set_title("Synthetic")
    book.set_language("en")
    chapters = []
    for ident, text in [("first", "First body"), ("second", "Second body")]:
        chapter = epub.EpubHtml(title="Same", file_name=ident + ".xhtml", uid=ident)
        chapter.content = f"<h1>Same</h1><p>{text}</p>"
        book.add_item(chapter)
        chapters.append(chapter)
    book.toc = tuple(chapters)
    book.spine = chapters
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    path = tmp_path / "spines.epub"
    epub.write_epub(str(path), book)
    rows = parse_document(path, "application/epub+zip")
    assert {r["spine"] for r in rows} == {"first", "second"}
    assert len({(r["spine"], r["paragraph"]) for r in rows}) == len(rows)
    assert all(r["chapter"] == "Same" and r["file_page"] is None for r in rows)
    assert all(
        r["parser_name"] == "ebooklib"
        and r["parser_version"] == importlib.metadata.version("ebooklib")
        for r in rows
    )
    assert len(list(_chunks([{**r, "id": uuid4()} for r in rows]))) == 2


def test_ingestion_metadata_readback_legacy_null_and_access(world, runtime_db):
    source_id = make_source(world, "Synthetic line one.\nSynthetic line two.")
    with psycopg.connect(world["dsn"]) as conn:
        job = submit_processing(conn, world["principal"], source_id, str(uuid4()), extract=False)
    run_job(world["worker"], job["job_id"], world["principal"], world["store"].root)
    with psycopg.connect(world["dsn"]) as conn:
        result = content.read_source(conn, world["principal"], source_id)
    fragments = result["fragments"]
    assert len(fragments) == 2
    assert all(
        r["parser_name"] == "python" and r["parser_version"] == platform.python_version()
        for r in fragments
    )
    assert all(r["heading"] is None and r["spine"] is None for r in fragments)
    run_job(world["worker"], job["job_id"], world["principal"], world["store"].root)
    with psycopg.connect(world["dsn"]) as conn:
        assert content.read_source(conn, world["principal"], source_id)["fragments"] == fragments
        with pytest.raises(AccessDenied):
            content.read_source(conn, world["other"], source_id)

    legacy = make_source(world, "Legacy original")
    with psycopg.connect(runtime_db) as conn:
        conn.execute(
            "INSERT INTO kb.fragment(source_id,ordinal,text,locator_kind) "
            "VALUES(%s,0,'Legacy text','none')",
            (legacy,),
        )
    with psycopg.connect(world["dsn"]) as conn:
        fragment = content.read_source(conn, world["principal"], legacy)["fragments"][0]
    assert all(
        fragment[key] is None for key in ("parser_name", "parser_version", "spine", "heading")
    )
